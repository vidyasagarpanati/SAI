"""Read a posturography ("body sway") export into typed trials.

IN     a .csv or .xlsx export, one row per trial, units carried in the header
DO     check the column contract, select one athlete, drop measures the
       instrument did not record, group trials by condition
OUT    Trial records and a per-measure availability map

Three rules come from the export that was supplied, and each one exists
because getting it wrong would put a false number in an evidence-based report.

1. ONE FILE HOLDS SEVERAL ATHLETES. The sample carried four. Selecting the
   wrong rows is silent and total, so a name that matches none or more than one
   is a hard error that names the candidates.
2. A COLUMN OF EXACT ZEROS IS NOT A MEASUREMENT. In the sample every
   contralateral and sample-entropy column was 0.000 in all twenty rows, which
   means a single-plate capture, not perfect left/right symmetry. Such a
   measure is reported unavailable, never as zero.
3. CONDITIONS ARE NOT COMPARABLE. Trials differ in task, stance and duration;
   a 50 s free stance and a 10 s draw-and-hold are different measurements. The
   condition is part of every key, so no table can average across them.
"""
from __future__ import annotations

import csv
import hashlib
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

# ------------------------------------------------------------------ contract
META_COLUMNS = {
    "project": "Project",
    "visit": "Visit",
    "date": "Date",
    "upper": "Upper Extremities",
    "sensory": "Sensory Manipulation",
    "stance": "Stance Position",
    "duration_s": "Duration [s]",
    "dual_task": "Dual Task",
    "footwear": "Footwear",
    "repetition": "Repetition",
    "subject_id": "Subject - ID",
    "subject_name": "Subject - Name",
    "subject_dob": "Subject - Date of birth",
    "subject_sex": "Subject - Sex",
    "subject_height_m": "Subject - Height [m]",
    "subject_weight_kg": "Subject - Weight [kg]",
    "subject_notes": "Subject - Notes",
}

_G = "Global parameters - General --- "
_I = "Global parameters - Interval specific --- "
_C = "Contralateral Parameters --- "

# Short key -> (units, exact column name). Deliberately a whitelist: the export
# carries 127 measures and the model must not be handed all of them.
MEASURES: dict[str, tuple[str, str]] = {
    "sway_path_total":      ("mm",     _G + "Sway path - total [mm]"),
    "sway_path_ap":         ("mm",     _G + "Sway path - A-P [mm]"),
    "sway_path_ml":         ("mm",     _G + "Sway path - M-L [mm]"),
    "sway_v_total":         ("mm/s",   _G + "Sway V - total [mm/s]"),
    "sway_v_ap":            ("mm/s",   _G + "Sway V - A-P [mm/s]"),
    "sway_v_ml":            ("mm/s",   _G + "Sway V - M-L [mm/s]"),
    "sway_amp_mean_ap":     ("mm",     _G + "Sway average amplitude - A-P [mm]"),
    "sway_amp_mean_ml":     ("mm",     _G + "Sway average amplitude - M-L [mm]"),
    "sway_amp_max_ap":      ("mm",     _G + "Sway maximal amplitude - A-P [mm]"),
    "sway_amp_max_ml":      ("mm",     _G + "Sway maximal amplitude - M-L [mm]"),
    "sway_area_total":      ("mm^2",   _G + "Sway area - total [mm^2]"),
    "ellipse_95":           ("mm^2",   _G + "Prediction ellipse area - 95% [mm^2]"),
    "fre_mean_ap":          ("Hz",     _G + "Mean FRE of total spectrum - A-P [Hz]"),
    "fre_mean_ml":          ("Hz",     _G + "Mean FRE of total spectrum - M-L [Hz]"),
    "sway_v_endurance":     ("%",      _I + "Sway V endurance index - total [%]"),
    "sway_v_fatigue":       ("%",      _I + "Sway V fatigue index - total [%]"),
    # Left/right loading. Whitelisted precisely so the all-zero rule can mark it
    # unavailable on a single-plate capture. Reporting these as 0 would read as
    # perfect symmetry, which is the worst failure an evidence-based report has.
    "sway_v_left":          ("mm/s",   _C + "Left leg - Sway V - total [mm/s]"),
    "sway_v_right":         ("mm/s",   _C + "Right leg - Sway V - total [mm/s]"),
    "lr_sway_v_ratio":      ("%",      _C + "Left/right leg - Sway V - total [%]"),
}

ALL_ZERO = "NOT RECORDED BY THE INSTRUMENT (column is zero in every trial)"
MISSING = "COLUMN ABSENT FROM THE EXPORT"


class ForcePlateError(ValueError):
    """The export cannot be read as evidence. Never downgraded to a warning."""


@dataclass
class Trial:
    condition: str                       # short label used in evidence keys
    condition_full: dict                 # every field that defines the condition
    repetition: int
    date: str
    values: dict[str, float] = field(default_factory=dict)


@dataclass
class Export:
    path: Path
    sha256: str
    athlete: str
    subject: dict
    trials: list[Trial]
    unavailable: dict[str, str]          # measure -> why
    other_athletes: list[str]
    n_rows_total: int


# ------------------------------------------------------------------ reading
def _rows_csv(path: Path) -> list[dict]:
    # csv.DictReader, never split(','): two columns are free-text notes and may
    # hold a comma. The supplied sample happens not to, so a naive split would
    # have passed today and broken on the next export.
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def _rows_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment problem, not logic
        raise ForcePlateError(
            f"Reading {path.name} needs openpyxl.  pip install openpyxl==3.1.5"
        ) from exc
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    it = sheet.iter_rows(values_only=True)
    try:
        header = ["" if c is None else str(c).strip() for c in next(it)]
    except StopIteration:
        raise ForcePlateError(f"{path.name} is empty.") from None
    rows = []
    for raw in it:
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue
        rows.append({h: ("" if v is None else v) for h, v in zip(header, raw)})
    book.close()
    return rows


def read_rows(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _rows_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _rows_xlsx(path)
    raise ForcePlateError(
        f"{path.name}: unsupported force-plate format '{suffix}'. "
        f"Export the trials as .csv or .xlsx."
    )


# ------------------------------------------------------------------ helpers
def _num(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return None if isinstance(value, float) and math.isnan(value) else float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def normalise_name(name: str) -> str:
    return re.sub(r"[^a-z]+", "", str(name or "").lower())


def _slug(text: str) -> str:
    return re.sub(r"_+", "_", re.sub(r"[^A-Z0-9]+", "_", str(text or "").upper())).strip("_")


def _condition(row: dict) -> tuple[str, str, dict]:
    """Return (short label, full label, the fields that define the condition)."""
    full = {key: str(row.get(META_COLUMNS[key], "")).strip()
            for key in ("upper", "sensory", "stance", "duration_s", "dual_task")}
    duration = _num(full["duration_s"])
    dur = f"{int(duration)}S" if duration is not None else "DURUNKNOWN"
    short = f"{_slug(full['upper'])}_{dur}"
    long = "_".join([short, _slug(full["sensory"]), _slug(full["stance"]),
                     _slug(full["dual_task"])])
    return short, long, full


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ------------------------------------------------------------------ main entry
def load(path: Path, athlete_name: str) -> Export:
    path = Path(path)
    rows = read_rows(path)
    if not rows:
        raise ForcePlateError(f"{path.name} has a header but no trial rows.")

    present = set(rows[0].keys())
    missing_meta = [col for col in META_COLUMNS.values() if col not in present]
    if missing_meta:
        raise ForcePlateError(
            f"{path.name} is not the expected posturography export.\n"
            f"  missing columns: {missing_meta}\n"
            f"  If the export template changed, update forceplate.META_COLUMNS "
            f"rather than letting the parser guess."
        )

    # A measure whose column is absent, or exactly zero in EVERY row of the
    # file, is not a measurement. Judged over the whole file, not one athlete,
    # because an instrument channel is either recording or it is not.
    unavailable: dict[str, str] = {}
    for key, (_units, column) in MEASURES.items():
        if column not in present:
            unavailable[key] = MISSING
            continue
        seen = [_num(r.get(column)) for r in rows]
        if all(v is not None and v == 0.0 for v in seen):
            unavailable[key] = ALL_ZERO

    wanted = normalise_name(athlete_name)
    by_name: dict[str, list[dict]] = {}
    for row in rows:
        by_name.setdefault(normalise_name(row.get(META_COLUMNS["subject_name"])), []).append(row)

    matches = [n for n in by_name if n and (n == wanted or wanted in n or n in wanted)]
    exact = [n for n in matches if n == wanted]
    if exact:
        matches = exact
    if not matches:
        names = sorted({str(r.get(META_COLUMNS["subject_name"], "")).strip() for r in rows})
        raise ForcePlateError(
            f"No trials for \"{athlete_name}\" in {path.name}.\n"
            f"  athletes in the file: {names}\n"
            f"  Fix --athlete, or export the right athlete's trials."
        )
    if len(matches) > 1:
        shown = sorted({str(r.get(META_COLUMNS["subject_name"], "")).strip()
                        for n in matches for r in by_name[n]})
        raise ForcePlateError(
            f"\"{athlete_name}\" matches more than one athlete in {path.name}: {shown}\n"
            f"  Pass the full name so the right rows are selected. "
            f"Guessing here would silently report another athlete's posture."
        )

    mine = by_name[matches[0]]
    first = mine[0]
    subject = {
        "name": str(first.get(META_COLUMNS["subject_name"], "")).strip(),
        "sex": str(first.get(META_COLUMNS["subject_sex"], "")).strip() or None,
        "height_cm": (lambda m: round(m * 100.0, 1) if m else None)(
            _num(first.get(META_COLUMNS["subject_height_m"]))),
        "weight_kg": _num(first.get(META_COLUMNS["subject_weight_kg"])),
        "notes": str(first.get(META_COLUMNS["subject_notes"], "")).strip() or None,
        # Date of birth is read but never carried out of this function: the
        # report prints the name and the operator-entered age, not a DOB.
    }

    # Two different full conditions must never collapse onto one short label.
    seen_labels: dict[str, str] = {}
    collides = False
    for row in mine:
        short, long, _ = _condition(row)
        if seen_labels.setdefault(short, long) != long:
            collides = True

    trials: list[Trial] = []
    for row in mine:
        short, long, full = _condition(row)
        values = {}
        for key, (_units, column) in MEASURES.items():
            if key in unavailable:
                continue
            value = _num(row.get(column))
            if value is not None:
                values[key] = value
        trials.append(Trial(
            condition=long if collides else short,
            condition_full=full,
            repetition=int(_num(row.get(META_COLUMNS["repetition"])) or 0),
            date=str(row.get(META_COLUMNS["date"], "")).strip(),
            values=values,
        ))

    others = sorted({str(r.get(META_COLUMNS["subject_name"], "")).strip()
                     for r in rows} - {subject["name"]})
    return Export(path=path, sha256=sha256_file(path), athlete=subject["name"],
                  subject=subject, trials=trials, unavailable=unavailable,
                  other_athletes=[o for o in others if o], n_rows_total=len(rows))
